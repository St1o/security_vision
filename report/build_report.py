"""Сборка отчёта по практике в формате .docx по образцу (рамка и штамп берутся из файла-шаблона).

    python report/build_report.py --template Полуян.docx --out report/Отчёт.docx [--start-page 8]

Тексты глав лежат в report/content/*.md, источники — в report/sources.py, рисунки — в report/figures/
и report/inputs/ (сюда кладут скриншоты и файлы из ppe_result.zip). Если рисунка нет, вместо него
вставляется жёлтая рамка-заглушка с подсказкой, что именно сюда вставить.
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from sources import SOURCES  # noqa: E402

FONT = "Times New Roman"
BODY_SZ = 28          # 14 пт
TABLE_SZ = 24         # 12 пт
LISTING_SZ = 18       # 9 пт
DESIGNATION = "ПП.990000.000 ПЗ"


# ---------------------------------------------------------------- разбор содержимого ----------

class Block:
    def __init__(self, kind, text="", **kw):
        self.kind, self.text, self.kw = kind, text, kw


def parse_content(files: list[Path]) -> list[Block]:
    blocks: list[Block] = []
    for path in files:
        lines = path.read_text(encoding="utf-8").splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            i += 1
            if not line.strip():
                continue
            if line.startswith("#### "):
                blocks.append(Block("h3", line[5:].strip()))
            elif line.startswith("### "):
                blocks.append(Block("h2", line[4:].strip()))
            elif line.startswith("## "):
                blocks.append(Block("h1", line[3:].strip()))
            elif line.startswith("# "):
                blocks.append(Block("h0", line[2:].strip()))
            elif line.startswith("- "):
                blocks.append(Block("bullet", line[2:].strip()))
            elif line.startswith("{table "):
                parts = [p.strip() for p in line.strip("{}").split(" | ")]
                key = parts[0].split()[1]
                widths = None
                for extra in parts[2:]:
                    if extra.startswith("w="):
                        widths = [float(x) for x in extra[2:].split(",")]
                rows = []
                while i < len(lines) and lines[i].startswith("|"):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                    i += 1
                blocks.append(Block("table", parts[1], key=key, widths=widths, rows=rows))
            elif line.startswith("{figure "):
                parts = [p.strip() for p in line.strip("{}").split(" | ")]
                key = parts[0].split()[1]
                blocks.append(Block("figure", parts[1], key=key, file=parts[2],
                                    width=float(parts[3]) if len(parts) > 3 else 15.0,
                                    hint=parts[4] if len(parts) > 4 else ""))
            elif line.startswith("{listing "):
                parts = [p.strip() for p in line.strip("{}").split(" | ")]
                blocks.append(Block("listing", parts[1], file=parts[0].split("file=")[1].strip()))
            elif line.startswith("{center "):
                blocks.append(Block("center", line[8:].rstrip("}").strip()))
            elif line.strip() == "{signature}":
                blocks.append(Block("signature"))
            else:
                blocks.append(Block("p", line.strip()))
    return blocks


def number_things(blocks: list[Block]):
    figs, tabs, cites = {}, {}, {}
    for b in blocks:
        if b.kind == "figure":
            figs[b.kw["key"]] = len(figs) + 1
        elif b.kind == "table":
            tabs[b.kw["key"]] = len(tabs) + 1

    def scan(text):
        for m in re.finditer(r"\[((?:@\w+\s*)+)\]", text):
            for key in re.findall(r"@(\w+)", m.group(1)):
                if key not in SOURCES:
                    raise KeyError(f"Нет источника «{key}» в sources.py")
                cites.setdefault(key, len(cites) + 1)

    for b in blocks:
        scan(b.text)
        for row in b.kw.get("rows", []):
            for cell in row:
                scan(cell)
    return figs, tabs, cites


def resolve(text: str, figs, tabs, cites) -> str:
    text = re.sub(r"@fig:(\w+)", lambda m: str(figs[m.group(1)]), text)
    text = re.sub(r"@tab:(\w+)", lambda m: str(tabs[m.group(1)]), text)

    def cite(m):
        nums = sorted(cites[k] for k in re.findall(r"@(\w+)", m.group(1)))
        return "[" + ", ".join(str(n) for n in nums) + "]"

    return re.sub(r"\[((?:@\w+\s*)+)\]", cite, text)


# ---------------------------------------------------------------- построение XML ---------------

def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def set_run_fmt(run, size, bold=False, font=FONT, highlight=False, italic=False):
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = _el("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + a), font)
    for tag in ("w:b", "w:bCs", "w:i", "w:sz", "w:szCs", "w:highlight"):
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    if bold:
        rpr.append(_el("w:b"))
        rpr.append(_el("w:bCs"))
    if italic:
        rpr.append(_el("w:i"))
    rpr.append(_el("w:sz", w_val=size))
    rpr.append(_el("w:szCs", w_val=size))
    if highlight:
        rpr.append(_el("w:highlight", w_val="yellow"))


def add_runs(par, text: str, size: int, bold=False, font=FONT):
    """Разбирает **жирный** и {{заглушки}} и добавляет прогоны в абзац."""
    for token in re.split(r"(\*\*.+?\*\*|\{\{.+?\}\})", text):
        if not token:
            continue
        if token.startswith("**"):
            r = par.add_run(token[2:-2])
            set_run_fmt(r, size, bold=True, font=font)
        elif token.startswith("{{"):
            r = par.add_run("[" + token[2:-2] + "]")
            set_run_fmt(r, size, bold=bold, font=font, highlight=True)
        else:
            r = par.add_run(token)
            set_run_fmt(r, size, bold=bold, font=font)


def ppr(par):
    return par._p.get_or_add_pPr()


def fmt_par(par, *, style_id=None, jc=None, first=None, left=None, line=360, before=0, after=0,
            keep_next=False, page_break=False, no_num=False, contextual=False, keep_lines=False):
    p = ppr(par)
    for child in list(p):
        p.remove(child)
    if style_id:
        p.append(_el("w:pStyle", w_val=style_id))
    if keep_next:
        p.append(_el("w:keepNext"))
    if keep_lines:
        p.append(_el("w:keepLines"))
    if page_break:
        p.append(_el("w:pageBreakBefore"))
    if no_num:
        num = _el("w:numPr")
        num.append(_el("w:ilvl", w_val=0))
        num.append(_el("w:numId", w_val=0))
        p.append(num)
    p.append(_el("w:spacing", w_before=before, w_after=after, w_line=line, w_lineRule="auto"))
    ind = {}
    if left is not None:
        ind["w_left"] = left
    if first is not None:
        ind["w_firstLine"] = first
    if ind:
        p.append(_el("w:ind", **ind))
    if contextual:
        p.append(_el("w:contextualSpacing"))
    if jc:
        p.append(_el("w:jc", w_val=jc))


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(_el("w:shd", w_val="clear", w_color="auto", w_fill=fill))


def set_borders(table):
    tblpr = table._tbl.tblPr
    borders = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el("w:" + side, w_val="single", w_sz=4, w_space=0, w_color="000000"))
    tblpr.append(borders)
    tblpr.append(_el("w:tblLayout", w_type="fixed"))
    tblpr.append(_el("w:tblCellMar"))
    mar = tblpr.find(qn("w:tblCellMar"))
    for side, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        mar.append(_el("w:" + side, w_w=w, w_type="dxa"))



PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def _sort_children(el, order):
    def key(child):
        name = child.tag.split("}")[1]
        return order.index(name) if name in order else len(order)
    kids = sorted(list(el), key=key)
    for child in list(el):
        el.remove(child)
    for child in kids:
        el.append(child)


def normalize(doc):
    """Порядок дочерних элементов по схеме OOXML (Word строже LibreOffice)."""
    root = doc.element.body
    for tag, order in (("w:pPr", PPR_ORDER), ("w:tcPr", TCPR_ORDER), ("w:tblPr", TBLPR_ORDER), ("w:rPr", RPR_ORDER)):
        for el in root.iter(qn(tag)):
            _sort_children(el, order)


# ---------------------------------------------------------------- сборщик ---------------------

class Builder:
    def __init__(self, template: Path, designation: str, start_page: int, toc: dict | None):
        self.doc = Document(str(template))
        self.designation = designation
        self.start_page = start_page
        self.toc_pages = toc or {}
        self.styles = {s.name.lower(): s.style_id for s in self.doc.styles}
        self.headings: list[tuple[int, str]] = []   # (уровень, текст) для оглавления

        # рамка и штамп: берём секцию основного текста (вторая в образце)
        sect_main = copy.deepcopy(self.doc.sections[1]._sectPr)
        footer = self.doc.sections[1].footer
        for t in footer._element.iter(qn("w:t")):
            if t.text and t.text.strip().startswith("ПП."):
                t.text = designation
        body = self.doc.element.body
        for child in list(body):
            body.remove(child)
        for old in sect_main.findall(qn("w:pgNumType")):
            sect_main.remove(old)
        pg = _el("w:pgNumType", w_start=start_page)
        # pgNumType должен идти после pgMar
        mar = sect_main.find(qn("w:pgMar"))
        mar.addnext(pg) if mar is not None else sect_main.append(pg)
        body.append(sect_main)

    # --- базовые абзацы
    def p(self, text, **kw):
        par = self.doc.add_paragraph()
        fmt = dict(jc="both", first=709, contextual=True)
        fmt.update(kw.pop("fmt", {}))
        fmt_par(par, **fmt)
        add_runs(par, text, kw.get("size", BODY_SZ), bold=kw.get("bold", False))
        return par

    def h0(self, text, toc=True):
        par = self.doc.add_paragraph()
        fmt_par(par, style_id=self.styles["heading 1"] if toc else None, jc="center", after=240, keep_next=True,
                keep_lines=True, page_break=True, no_num=True)
        add_runs(par, text, BODY_SZ, bold=True)
        if toc:
            self.headings.append((1, text))

    def h1(self, text):
        par = self.doc.add_paragraph()
        fmt_par(par, style_id=self.styles["heading 1"], jc="left", first=709, after=240, keep_next=True,
                keep_lines=True, page_break=True, no_num=True)
        add_runs(par, text, BODY_SZ, bold=True)
        self.headings.append((1, text))

    def h2(self, text):
        par = self.doc.add_paragraph()
        fmt_par(par, style_id=self.styles["heading 2"], jc="left", first=709, before=120, after=240,
                keep_next=True, keep_lines=True, no_num=True)
        add_runs(par, text, BODY_SZ, bold=True)
        self.headings.append((2, text))

    def h3(self, text):
        par = self.doc.add_paragraph()
        fmt_par(par, jc="left", first=709, before=240, after=120, keep_next=True, keep_lines=True)
        add_runs(par, text, BODY_SZ, bold=True)

    def center(self, text):
        par = self.doc.add_paragraph()
        fmt_par(par, jc="center", after=240, keep_next=True)
        add_runs(par, text, BODY_SZ, bold=True)

    def bullet(self, text):
        par = self.doc.add_paragraph()
        fmt_par(par, jc="both", left=1134, contextual=True)
        pr = ppr(par)
        num = _el("w:numPr")
        num.append(_el("w:ilvl", w_val=0))
        num.append(_el("w:numId", w_val=9))
        pr.insert(0, num)
        pr.insert(1, _el("w:tabs"))
        pr.find(qn("w:tabs")).append(_el("w:tab", w_val="left", w_pos=993))
        add_runs(par, text, BODY_SZ)

    def signature(self):
        prev = self.doc.paragraphs[-1]
        ppr(prev).insert(0, _el("w:keepNext"))
        self.p("", fmt=dict(first=0, keep_next=True))
        par = self.doc.add_paragraph()
        fmt_par(par, jc="left", first=0, keep_next=True)
        add_runs(par, "«___» ___________ 2026 г.\t\t\t______________", BODY_SZ)
        par2 = self.doc.add_paragraph()
        fmt_par(par2, jc="right", first=0, left=0)
        add_runs(par2, "(подпись)          ", 20)

    # --- таблицы, рисунки, листинги
    def table(self, number, caption, rows, widths, resolve_fn):
        cap = self.doc.add_paragraph()
        fmt_par(cap, jc="left", first=0, before=120, after=60, keep_next=True, line=360)
        add_runs(cap, f"Таблица {number} – {caption}", BODY_SZ)
        ncols = len(rows[0])
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        set_borders(t)
        widths = widths or [17.0 / ncols] * ncols
        t.autofit = False
        for ci, w in enumerate(widths):
            t.columns[ci].width = Cm(w)
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(_el("w:cantSplit"))
            if ri == 0:
                trpr.append(_el("w:tblHeader"))
            for ci in range(ncols):
                cell = tr.cells[ci]
                cell.width = Cm(widths[ci])
                par = cell.paragraphs[0]
                keep = len(rows) <= 12 and ri < len(rows) - 1   # небольшая таблица не разрывается по страницам
                fmt_par(par, jc="center" if ri == 0 else "left", first=0, line=240, keep_next=keep)
                text = resolve_fn(row[ci] if ci < len(row) else "")
                add_runs(par, text, TABLE_SZ, bold=(ri == 0))
                tcpr = cell._tc.get_or_add_tcPr()
                tcpr.append(_el("w:vAlign", w_val="center"))
                if ri == 0:
                    shade(cell, "E7E6E6")
        after = self.doc.add_paragraph()
        fmt_par(after, first=0, line=240, after=120)

    def figure(self, number, caption, file, width_cm, hint):
        path = next((p for p in (HERE / "inputs" / file, HERE / "figures" / file) if p.exists()), None)
        if path:
            par = self.doc.add_paragraph()
            fmt_par(par, jc="center", first=0, before=120, after=60, keep_next=True, line=240)
            par.add_run().add_picture(str(path), width=Cm(width_cm))
        else:
            t = self.doc.add_table(rows=1, cols=1)
            set_borders(t)
            t.autofit = False
            t.columns[0].width = Cm(15)
            cell = t.rows[0].cells[0]
            cell.width = Cm(15)
            shade(cell, "FFF2CC")
            trpr = t.rows[0]._tr.get_or_add_trPr()
            trpr.append(_el("w:trHeight", w_val=int(Cm(5.5).twips), w_hRule="atLeast"))
            trpr.append(_el("w:cantSplit"))
            cell._tc.get_or_add_tcPr().append(_el("w:vAlign", w_val="center"))
            par = cell.paragraphs[0]
            fmt_par(par, jc="center", first=0, line=240)
            add_runs(par, "{{ВСТАВИТЬ РИСУНОК: " + (hint or file) + "}}", TABLE_SZ)
            gap = self.doc.add_paragraph()
            fmt_par(gap, first=0, line=120, keep_next=True)
        cap = self.doc.add_paragraph()
        fmt_par(cap, jc="center", first=0, after=240, line=360)
        add_runs(cap, f"Рисунок {number} – {caption}", BODY_SZ)

    def listing(self, caption, relpath):
        title = self.doc.add_paragraph()
        fmt_par(title, jc="left", first=709, before=120, after=120, keep_next=True, line=360)
        add_runs(title, caption, BODY_SZ)
        for raw in (ROOT / relpath).read_text(encoding="utf-8").expandtabs(4).splitlines():
            par = self.doc.add_paragraph()
            fmt_par(par, jc="left", first=0, line=240)
            r = par.add_run(raw if raw else " ")
            set_run_fmt(r, LISTING_SZ, font="Courier New")
            r._r.find(qn("w:t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    def sources_list(self, cites):
        for key, num in sorted(cites.items(), key=lambda kv: kv[1]):
            par = self.doc.add_paragraph()
            fmt_par(par, jc="left", first=709, line=360, contextual=True)  # слева: длинные ссылки при выравнивании по ширине рвут строки
            add_runs(par, f"{num}. {SOURCES[key]}", BODY_SZ)

    # --- оглавление
    def toc(self):
        self.h0("Содержание", toc=False)
        entries = [(lvl, text) for lvl, text in self.toc_entries]
        first = True
        for idx, (lvl, text) in enumerate(entries):
            par = self.doc.add_paragraph()
            style = self.styles.get("toc 1" if lvl == 1 else "toc 2")
            fmt_par(par, style_id=style, jc="left", first=0, left=0 if lvl == 1 else 567, line=360)
            ppr(par).append(_el("w:tabs"))
            ppr(par).find(qn("w:tabs")).append(_el("w:tab", w_val="right", w_leader="dot", w_pos=9500))
            if first:  # начало поля оглавления
                r = par.add_run()
                r._r.append(_el("w:fldChar", w_fldCharType="begin"))
                r2 = par.add_run()
                it = _el("w:instrText")
                it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                it.text = ' TOC \\o "1-2" \\h \\z \\u '
                r2._r.append(it)
                r3 = par.add_run()
                r3._r.append(_el("w:fldChar", w_fldCharType="separate"))
                first = False
            page = self.toc_pages.get(text, "")
            add_runs(par, text, BODY_SZ if lvl == 1 else 26)
            r = par.add_run("\t" + (str(page) if page != "" else ""))
            set_run_fmt(r, BODY_SZ if lvl == 1 else 26)
            if idx == len(entries) - 1:
                r = par.add_run()
                r._r.append(_el("w:fldChar", w_fldCharType="end"))


def build(template: Path, out: Path, start_page: int, toc_pages: dict | None, designation: str):
    files = sorted((HERE / "content").glob("*.md"))
    blocks = parse_content(files)
    figs, tabs, cites = number_things(blocks)
    res = lambda s: resolve(s, figs, tabs, cites)  # noqa: E731

    b = Builder(template, designation, start_page, toc_pages)
    b.toc_entries = [(1, "Введение")]
    for blk in blocks:
        if blk.kind == "h1":
            b.toc_entries.append((1, blk.text))
        elif blk.kind == "h2":
            b.toc_entries.append((2, blk.text))
        elif blk.kind == "h0" and blk.text != "Введение":
            b.toc_entries.append((1, blk.text))
    # перечень ресурсов вставляется перед приложениями
    sources_title = "Перечень использованных информационных ресурсов"
    idx_app = next(i for i, e in enumerate(b.toc_entries) if e[1].startswith("Приложение"))
    b.toc_entries.insert(idx_app, (1, sources_title))

    b.toc()
    sources_done = False
    for blk in blocks:
        if blk.kind == "h0" and blk.text.startswith("Приложение") and not sources_done:
            b.h0(sources_title)
            b.sources_list(cites)
            sources_done = True
        k = blk.kind
        if k == "h0":
            b.h0(blk.text)
        elif k == "h1":
            b.h1(blk.text)
        elif k == "h2":
            b.h2(blk.text)
        elif k == "h3":
            b.h3(blk.text)
        elif k == "center":
            b.center(blk.text)
        elif k == "p":
            b.p(res(blk.text))
        elif k == "bullet":
            b.bullet(res(blk.text))
        elif k == "table":
            b.table(tabs[blk.kw["key"]], blk.text, blk.kw["rows"], blk.kw["widths"], res)
        elif k == "figure":
            b.figure(figs[blk.kw["key"]], blk.text, blk.kw["file"], blk.kw["width"], blk.kw["hint"])
        elif k == "listing":
            b.listing(blk.text, blk.kw["file"])
        elif k == "signature":
            b.signature()
    if not sources_done:
        b.h0(sources_title)
        b.sources_list(cites)

    # убрать из пакета картинки, на которые больше нет ссылок (сканы титульных листов образца)
    part = b.doc.part
    xml = part._element.xml if hasattr(part, "_element") else ""
    for rid, rel in list(part.rels.items()):
        if rel.reltype.endswith("/image") and f'"{rid}"' not in xml:
            part.drop_rel(rid)
    normalize(b.doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    b.doc.save(str(out))
    return b


def locate_pages(docx_path: Path, entries, start_page: int) -> dict:
    """Конвертирует в PDF через LibreOffice и находит страницу каждого заголовка (для оглавления)."""
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(docx_path)],
                   check=True, capture_output=True, timeout=300)
    pdf = tmp / (docx_path.stem + ".pdf")
    pages = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout.split("\f")
    result, cursor = {}, 1  # страница 0 — оглавление
    norm = lambda s: re.sub(r"\s+", " ", s).strip()  # noqa: E731
    for lvl, text in entries:
        key = norm(text)[:45]
        for pi in range(cursor, len(pages)):
            body = norm(pages[pi])
            if key in body:
                result[text] = start_page + pi
                cursor = pi
                break
    return result


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--start-page", type=int, default=8)
    ap.add_argument("--designation", default=DESIGNATION)
    ap.add_argument("--no-toc-pages", action="store_true", help="не рассчитывать номера страниц оглавления")
    a = ap.parse_args()

    first = build(a.template, a.out, a.start_page, None, a.designation)
    if not a.no_toc_pages:
        pages = locate_pages(a.out, first.toc_entries, a.start_page)
        build(a.template, a.out, a.start_page, pages, a.designation)
        missing = [t for _, t in first.toc_entries if t not in pages]
        print("страниц оглавления найдено:", len(pages), "из", len(first.toc_entries), "| не найдено:", missing)
    print("готово:", a.out)

import json
from datetime import date, datetime
from pathlib import Path

import pytest

from svision.config import ConfigError, load_config
from svision.geometry import iou, point_in_polygon
from svision.report import build_act
from svision.shifts import Shift, shift_for, validate_shifts
from svision.storage import Storage
from svision.tracking import IoUTracker, TrackerParams

ROOT = Path(__file__).resolve().parents[1]
SHIFTS = (Shift(1, "Смена 1", 0, 480), Shift(2, "Смена 2", 480, 960), Shift(3, "Смена 3", 960, 1440))


def test_geometry():
    assert iou((0, 0, 10, 10), (0, 0, 10, 10)) == 1.0
    assert iou((0, 0, 10, 10), (20, 20, 30, 30)) == 0.0
    square = [(0, 0), (1, 0), (1, 1), (0, 1)]
    assert point_in_polygon(0.5, 0.5, square)
    assert not point_in_polygon(1.5, 0.5, square)


def test_tracker_keeps_id_and_drops_lost_tracks():
    tracker = IoUTracker(TrackerParams(iou_threshold=0.3, max_missed=2, min_hits=2))
    assert tracker.update([(0, 0, 100, 200)]) == []  # новый трек подтверждается со 2-го кадра
    out = tracker.update([(5, 0, 105, 200)])
    assert [i for i, _ in out] == [1]
    out = tracker.update([(10, 0, 110, 200)])
    assert [i for i, _ in out] == [1]
    for _ in range(3):
        tracker.update([])
    tracker.update([(10, 0, 110, 200)])
    out = tracker.update([(12, 0, 112, 200)])
    assert [i for i, _ in out] == [2]  # старый трек удалён, новому выдан новый номер


def test_shifts():
    validate_shifts(SHIFTS)
    assert shift_for(datetime(2026, 10, 20, 7, 59), SHIFTS).number == 1
    assert shift_for(datetime(2026, 10, 20, 8, 0), SHIFTS).number == 2
    assert shift_for(datetime(2026, 10, 20, 23, 59), SHIFTS).number == 3
    with pytest.raises(ValueError):
        validate_shifts((Shift(1, "a", 0, 480), Shift(2, "b", 600, 1440)))


def test_storage_roundtrip(tmp_path):
    db = Storage(tmp_path / "v.db")
    db.init()
    rid = db.add(datetime(2026, 10, 20, 9, 15, 3), 2, "Камера 1", "Зона А", "no_helmet", track_id=4)
    db.add(datetime(2026, 10, 20, 17, 0, 0), 3, "Камера 1", "Зона А", "no_vest", is_demo=True)
    db.add(datetime(2026, 11, 2, 1, 0, 0), 1, "Камера 1", "Зона А", "no_vest")
    db.set_duration(rid, 12.5)
    assert db.days_with_violations(2026, 10) == {date(2026, 10, 20)}
    assert db.counts_by_shift(date(2026, 10, 20)) == {2: 1, 3: 1}
    rec = db.list(date(2026, 10, 20), 2)[0]
    assert (rec.kind, rec.track_id, rec.duration_s, rec.label) == ("no_helmet", 4, 12.5, "Нет каски")
    assert db.delete_demo() == 1
    assert db.counts_by_shift(date(2026, 10, 20)) == {2: 1}


def test_default_config_loads():
    cfg = load_config(ROOT / "config" / "default.json")
    assert cfg.required_roles == {"helmet", "vest"}
    assert [s.number for s in cfg.shifts] == [1, 2, 3]
    assert cfg.model.path == ROOT / "models" / "ppe.pt"


def test_bad_config_gives_readable_error(tmp_path):
    bad = tmp_path / "c.json"
    raw = json.loads((ROOT / "config" / "default.json").read_text(encoding="utf-8"))
    raw["zones"] = [{"name": "x", "polygon": [[0, 0], [1, 1]], "required": ["helmet"]}]
    bad.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ConfigError, match="минимум 3 точки"):
        load_config(bad, root=tmp_path)


def test_act_is_built(tmp_path):
    db = Storage(tmp_path / "v.db")
    db.init()
    db.add(datetime(2026, 10, 20, 9, 15, 3), 2, "Камера 1", "Зона А", "no_helmet", track_id=4, duration_s=12, is_demo=True)
    records = db.list(date(2026, 10, 20), 2)
    out = build_act(tmp_path / "act.docx", records, date(2026, 10, 20), SHIFTS[1], "Камера 1",
                    "Объект", "норматив", db.snapshot_path)
    from docx import Document

    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "АКТ № 20261020-2" in text and "ТЕСТОВЫЕ" in text
    table = Document(str(out)).tables[0]
    assert table.rows[1].cells[3].text == "Нет каски"

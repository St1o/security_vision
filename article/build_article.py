"""Сборка статьи в .docx по образцу (поля, шрифт 14 пт, абзацный отступ 1,25 см, шапка, аннотации на двух языках, таблицы).

    python article/build_article.py --template образец.docx --out article/out/Статья.docx

Страница, стили и разметка берутся из файла-образца; тексты статьи заданы в этом файле.
Источники нумеруются автоматически по порядку первого упоминания: [@ключ] или [@ключ1,@ключ2].
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "report"))
from sources import SOURCES  # noqa: E402

FONT = "Times New Roman"
ACCESS_DATE = "22.10.2026"
EMAIL = "gridasov.a@internet.ru"
INDENT = 709  # 1,25 см


# ------------------------------------------------------------------ текст статьи ----------------

UDK = "УДК 004.93"

AUTHOR_RU = [
    "А.Ю. Гридасов, обучающийся,",
    "специальности «10.05.01»,",
    "ФГБОУ ВО «Донской государственный технический университет»,",
    "EMAIL",
    "г. Ростов-на-Дону, Россия",
]
TITLE_RU = ("Контроль применения средств индивидуальной защиты на строительной площадке "
            "методами компьютерного зрения: подтверждение нарушений по окну кадров")
ANNOT_RU = (
    "Статья посвящена автоматизации контроля применения средств индивидуальной защиты (СИЗ) на строительной "
    "площадке методами компьютерного зрения. Рассмотрены требования охраны труда к применению каски и сигнального "
    "жилета, а также существующие решения на рынке. Разработан прототип системы, который по видеозаписи обнаруживает "
    "людей, каски и жилеты нейросетью YOLO11n, присваивает людям номера, привязывает СИЗ к человеку по взаимному "
    "расположению рамок и подтверждает нарушение по доле кадров без СИЗ в окне длиной 3 секунды. На тестовой выборке "
    "набора Construction Site Safety модель достигает mAP50 0,842 для каски и 0,805 для жилета. Моделирование показало "
    "отсутствие ложных нарушений у работника в каске за минуту наблюдения (0 из 1000 прогонов) и фиксацию настоящего "
    "нарушения за 6 секунд в 98 % случаев. Прототип работает на ноутбуке без видеокарты со скоростью около 23 кадров в "
    "секунду, ведёт журнал нарушений по рабочим сменам и формирует акт, не определяя личность работника."
)
KEYWORDS_RU = ("компьютерное зрение, средства индивидуальной защиты, охрана труда, обнаружение объектов, YOLO, "
               "строительная площадка.")

AUTHOR_EN = [
    "Gridasov A.Yu., student,",
    "specialty 10.05.01,",
    "Don State Technical University,",
    "Rostov-on-Don, Russia",
]
TITLE_EN = ("Monitoring the Use of Personal Protective Equipment on a Construction Site by Computer Vision Methods: "
            "Confirming Violations over a Frame Window")
ANNOT_EN = (
    "The article is devoted to automating the monitoring of personal protective equipment (PPE) use at a construction "
    "site by computer vision methods. The occupational safety requirements for the use of hard hats and high-visibility "
    "vests and the existing solutions on the market are considered. A prototype system has been developed that detects "
    "people, hard hats and vests in a video with the YOLO11n neural network, assigns numbers to people, associates PPE "
    "with a person by the mutual position of bounding boxes, and confirms a violation by the share of frames without "
    "PPE within a 3-second window. On the test set of the Construction Site Safety dataset the model achieves mAP50 of "
    "0.842 for hard hats and 0.805 for vests. Simulation showed no false violations for a worker wearing a hard hat "
    "during one minute of observation (0 out of 1000 runs) and detection of a real violation within 6 seconds in 98% of "
    "cases. The prototype runs on a laptop without a graphics card at about 23 frames per second, keeps a violation log "
    "by work shifts and generates a violation report without identifying the worker."
)
KEYWORDS_EN = "computer vision, personal protective equipment, occupational safety, object detection, YOLO, construction site."

# разделы: ("h", заголовок) | ("p", абзац) | ("table", ключ) | ("gap",)
BODY = [
    ("h", "Введение"),
    ("p", "Строительная площадка относится к объектам с повышенной опасностью: на ней одновременно работают люди, "
          "строительная техника и перемещаемые грузы. Наиболее массовыми средствами индивидуальной защиты (СИЗ) остаются "
          "защитная каска и сигнальный жилет. Работодатель обязан обеспечить работников СИЗ, а работник – правильно "
          "применять выданные СИЗ [@tk]; отраслевые требования к безопасности работ при строительстве установлены "
          "Правилами по охране труда, утверждёнными приказом Минтруда России от 11.12.2020 № 883н [@pot883]. Однако на "
          "практике контроль применения СИЗ остаётся выборочным: обход ответственного лица охватывает ограниченную часть "
          "площадки и рабочего времени, а оператор системы видеонаблюдения теряет концентрацию при просмотре однообразного "
          "видеоряда уже через десятки минут [@antimalware]. Методы компьютерного зрения позволяют автоматизировать эту "
          "рутинную проверку и фиксировать нарушения объективно: со временем, местом и фотографией."),
    ("p", "Целью работы является разработка и испытание прототипа системы компьютерного зрения, которая по видеозаписи "
          "обнаруживает отсутствие каски и сигнального жилета у работников и ведёт журнал нарушений по рабочим сменам. "
          "Для достижения цели необходимо рассмотреть требования к применению СИЗ и существующие решения, выбрать и "
          "дообучить модель обнаружения объектов, разработать алгоритм привязки СИЗ к человеку и подтверждения нарушения, "
          "а также реализовать программу и оценить её работу."),

    ("h", "Существующие решения и их ограничения"),
    ("p", "Рынок систем контроля СИЗ представлен универсальными платформами видеонаблюдения, специализированными "
          "облачными сервисами охраны труда и разработками отдельных предприятий. Например, в платформе Axxon One "
          "детектор СИЗ распознаёт каску, капюшон, перчатки и жилет; в его основе две нейросети: первая разделяет тело "
          "человека на зоны (голова, плечи, руки, туловище), вторая классифицирует СИЗ на выбранной части тела. "
          "Производитель заявляет точность 80–95 % и 5–20 % ложных срабатываний для стандартной сети, не обученной на "
          "данных объекта, и рекомендует дообучение; для работы необходимы дополнительный модуль и лицензия [@axxon_docs]. "
          "Известны и внедрения на промышленных предприятиях, в том числе с дополнительной идентификацией работников по "
          "лицу [@antimalware]."),
    ("p", "Общими особенностями таких решений являются коммерческая лицензия, необходимость серверной инфраструктуры и "
          "интеграции с системой видеонаблюдения, а также зависимость точности от дообучения на данных объекта "
          "[@axxon_docs]. Идентификация работника по лицу относится к обработке биометрических персональных данных и, как "
          "правило, требует письменного согласия субъекта [@fz152]. Поэтому в разработанном прототипе личность работника не "
          "определяется: система фиксирует факт нарушения, время, место и фотографию с номером человека на видео, а "
          "установление личности остаётся за ответственным по охране труда."),

    ("h", "Модель обнаружения объектов"),
    ("p", "Задача обнаружения касок на видеозаписях с камер наблюдения рассматривалась в ряде работ. В [@fang] "
          "предложен метод глубокого обучения для обнаружения работников без касок на видео дальнего плана, а в [@wu] – "
          "метод обнаружения касок на строительной площадке и набор данных для сравнения таких методов. Для работы в "
          "реальном времени на обычном компьютере подходят одностадийные детекторы семейства YOLO, которые определяют "
          "положение и класс всех объектов за один проход нейросети [@yolo]. В прототипе применена наименьшая модель "
          "семейства YOLO11 – YOLO11n [@yolo11], предобученная на наборе COCO [@coco] и дообученная на открытом наборе "
          "Construction Site Safety [@rf_css]."),
    ("p", "Набор содержит 2801 изображение (2605 обучающих, 114 проверочных и 82 тестовых) и десять классов; системе "
          "необходимы классы Person, Hardhat и Safety Vest, а классы NO-Hardhat и NO-Safety Vest используются как "
          "дополнительное свидетельство. Обучение выполнено в облачной среде Google Colab на видеокарте T4: за 35 минут "
          "пройдено 62 эпохи, лучшие веса получены на эпохе 56 (обучение остановлено по лимиту времени, а не по "
          "критерию достижения плато). Результаты на тестовой выборке из 82 изображений приведены в таблице 1."),
    ("table", "metrics"),
    ("p", "Для каски и жилета mAP50 составляет 0,842 и 0,805. Полнота обнаружения каски (0,755) означает, что примерно "
          "в каждом четвёртом кадре каска на добросовестном работнике не находится, а классы отсутствия СИЗ "
          "распознаются хуже (для NO-Hardhat mAP50 равен 0,520). Это определяет архитектуру алгоритма: решение о "
          "нарушении нельзя принимать по одному кадру. Следует учитывать небольшой размер тестовой выборки: полученные "
          "оценки приблизительны, а недообученность модели оставляет резерв для улучшения."),

    ("h", "Алгоритм привязки СИЗ к человеку и подтверждения нарушений"),
    ("p", "Программа обрабатывает видеозапись покадрово. На каждом кадре нейросеть находит людей, каски и жилеты; "
          "трекер по пересечению рамок (IoU) присваивает каждому человеку номер, сохраняющийся от кадра к кадру; "
          "положение человека относительно зоны контроля определяется по точке опоры – середине нижней границы рамки."),
    ("p", "Нейросеть выдаёт независимые рамки, поэтому необходимо определить, какой каске и какому жилету принадлежит "
          "каждый человек. Каска засчитывается, если центр её рамки лежит в области головы: от 15 % высоты рамки "
          "человека выше её верхней границы до 40 % высоты ниже неё, по ширине область расширена на 15 %. Жилет "
          "засчитывается, если центр его рамки лежит в области торса (от 15 до 75 % высоты рамки) и не менее половины "
          "площади рамки жилета находится внутри рамки человека. Так каска, лежащая на земле или висящая на поясе, не "
          "засчитывается. Если используются классы отсутствия СИЗ, результат определяется сравнением уверенности "
          "нейросети для классов «есть» и «нет»."),
    ("p", "Результат отдельного кадра ненадёжен. При полноте обнаружения каски около 0,75 вероятность того, что за "
          "3 секунды (30 кадров при 10 кадрах в секунду) найдётся хотя бы один пропуск, составляет 1 – 0,75^{30} ≈ 0,9998, "
          "поэтому правило «нарушение, если каски нет три секунды подряд» признавало бы нарушителем практически любого "
          "работника в каске. В окончательном алгоритме нарушение подтверждается по доле кадров в окне: если окно длиной "
          "3 с покрыто наблюдениями не менее чем на 90 %, а доля кадров без СИЗ в нём составляет не менее 0,8. "
          "Нарушение считается устранённым, если в окне 2 с СИЗ присутствует не менее чем в половине кадров. Повторная "
          "запись у того же человека (или на том же месте при смене номера) в течение 30 с не создаётся."),
    ("p", "Устойчивость правила проверена моделированием методом Монте-Карло на настоящем коде программы: "
          "последовательности обнаружений с частотой 10 кадров в секунду, кадры считались независимыми, на каждое "
          "условие выполнено 1000 прогонов. Вероятность пропуска каски у работника в каске принималась равной 0,25 "
          "(по полноте) и 0,31 (по матрице ошибок), вероятность ошибочного обнаружения каски у работника без неё – 0,17 "
          "(по матрице ошибок). Результаты приведены в таблице 2."),
    ("table", "montecarlo"),
    ("p", "Правило практически не создаёт ложных нарушений у работников в касках, а за несколько секунд подтверждает "
          "настоящее нарушение даже тогда, когда модель в части кадров ошибочно «видит» каску; медианное время фиксации "
          "составило 2,7 с. Ограничением расчёта является допущение о независимости кадров: если каска закрыта "
          "предметом, она может отсутствовать в кадрах подряд, поэтому на настоящем видео ложные срабатывания возможны "
          "чаще."),

    ("h", "Программная реализация и результаты испытаний"),
    ("p", "Программа написана на языке Python: графический интерфейс реализован на PySide6, обработка видео – на "
          "OpenCV, журнал нарушений хранится в SQLite, акты формируются в формате Word библиотекой python-docx. Окно "
          "«Мониторинг» показывает видео с разметкой (зелёная рамка – СИЗ в норме, жёлтая – идёт проверка, красная – "
          "нарушение) и таблицу текущих нарушений. Вкладка «Журнал нарушений» содержит календарь с выделением дат, на "
          "которые были нарушения, три рабочие смены по 8 часов и таблицу с фотографией выбранного нарушения; по смене "
          "формируется акт. Логика проверена 28 автоматическими тестами."),
    ("p", "Испытания проведены на четырёх видеозаписях длительностью от 6 до 27 секунд, снятых на строительных "
          "площадках. Программа зафиксировала 10 нарушений (6 – отсутствие каски, 4 – отсутствие сигнального жилета) "
          "длительностью от 3 до 16 с. Количественная оценка верных и ложных срабатываний на этих роликах не "
          "проводилась, так как для неё требуется ручная разметка, а четыре короткие записи недостаточны для "
          "статистических выводов; поэтому каждая запись сопровождается фотографией для визуальной проверки "
          "ответственным лицом."),
    ("p", "Скорость проверена на ноутбуке с процессором Intel Core i3-1215U (8 ГБ ОЗУ, без видеокарты). Модель в "
          "формате PyTorch обрабатывает кадр 1280 × 720 за 44 мс (около 22,6 кадра в секунду), в самом приложении "
          "скорость составила около 23 кадров в секунду, что в 4,5 раза выше требуемых 5. Формат OpenVINO показал по "
          "медиане 21 мс, но среднее время на кадр оказалось равным 91 мс, а в приложении скорость упала до 4,1 кадра "
          "в секунду, поэтому основным оставлен PyTorch; выбор формата следует подтверждать замером на целевом "
          "компьютере."),
    ("p", "Система проверяет наличие каски и жилета, но не их исправность, сертификацию и правильность ношения; точность "
          "зависит от набора данных и ракурса камеры; простой трекер после долгого перекрытия может присвоить человеку "
          "новый номер. Решение о взыскании принимает ответственный по охране труда, а система остаётся "
          "вспомогательным средством."),

    ("h", "Заключение"),
    ("p", "Разработан и испытан прототип системы компьютерного зрения для контроля применения каски и сигнального "
          "жилета на строительной площадке. Модель YOLO11n, дообученная на открытом наборе, достигает mAP50 0,842 для "
          "каски и 0,805 для жилета; привязка СИЗ к человеку по положению рамок и подтверждение нарушения по доле "
          "кадров в окне 3 с позволяют не создавать ложных нарушений у работников в касках (0 из 1000 прогонов "
          "моделирования) и фиксировать настоящее нарушение за 6 секунд в 98 % случаев."),
    ("p", "Таким образом, цель работы – разработка и испытание прототипа системы контроля применения СИЗ по видео – "
          "достигнута. Прототип работает на ноутбуке без видеокарты в реальном времени, ведёт журнал по рабочим сменам и "
          "формирует акт, не определяя личность работника."),
    ("p", "Перспективными направлениями дальнейших работ являются дообучение модели на данных конкретного объекта и "
          "увеличение числа обучающих эпох, расширение перечня контролируемых СИЗ, применение более устойчивого "
          "трекера, количественная оценка на размеченных видеозаписях и, при наличии согласия работников, "
          "идентификация по номерам на жилетах."),
]

TABLES = {
    "metrics": {
        "caption": "Таблица 1 – Показатели модели YOLO11n на тестовой выборке набора Construction Site Safety",
        "widths": [3000, 1550, 1550, 1600, 1596],
        "rows": [
            ["Класс", "Precision", "Recall", "mAP50", "mAP50–95"],
            ["Все классы", "0,863", "0,648", "0,713", "0,433"],
            ["Hardhat (каска)", "0,961", "0,755", "0,842", "0,522"],
            ["Safety Vest (жилет)", "0,866", "0,738", "0,805", "0,501"],
            ["Person (человек)", "0,830", "0,724", "0,805", "0,480"],
            ["NO-Hardhat (нет каски)", "0,813", "0,531", "0,520", "0,289"],
            ["NO-Safety Vest (нет жилета)", "0,898", "0,683", "0,762", "0,468"],
        ],
    },
    "montecarlo": {
        "caption": "Таблица 2 – Результаты моделирования правила подтверждения нарушений (1000 прогонов на условие)",
        "widths": [5200, 4096],
        "rows": [
            ["Условие", "Результат"],
            ["Работник в каске, пропуск каски 0,25 (60 с)", "Ложных нарушений нет (0 из 1000)"],
            ["Работник в каске, пропуск каски 0,31 (60 с)", "Ложных нарушений нет (0 из 1000)"],
            ["Работник без каски: фиксация за 4 с", "86,9 % прогонов"],
            ["Работник без каски: фиксация за 6 с", "97,8 % прогонов"],
            ["Работник без каски: фиксация за 8 с", "99,6 % прогонов"],
            ["Работник без каски: фиксация за 12 с", "99,9 % прогонов"],
        ],
    },
}

# источники в стиле образца: «Фамилия, И. О. Название / И. О. Фамилия // Издание. – Год. – ...»
BIB = {
    "fang": "Fang, Q. Detecting non-hardhat-use by a deep learning method from far-field surveillance videos / Q. Fang, "
            "H. Li, X. Luo, L. Ding, H. Luo, T. M. Rose, W. An // Automation in Construction. – 2018. – Vol. 85. – "
            "P. 1-9.",
    "wu": "Wu, J. Automatic detection of hardhats worn by construction personnel: A deep learning approach and benchmark "
          "dataset / J. Wu, N. Cai, W. Chen, H. Wang, G. Wang // Automation in Construction. – 2019. – Vol. 106. – "
          "Art. 102894.",
    "yolo": "Redmon, J. You Only Look Once: Unified, Real-Time Object Detection / J. Redmon, S. Divvala, R. Girshick, "
            "A. Farhadi // Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition. – 2016. – "
            "P. 779-788.",
    "coco": "Lin, T.-Y. Microsoft COCO: Common Objects in Context / T.-Y. Lin, M. Maire, S. Belongie [et al.] // "
            "Computer Vision – ECCV 2014. – 2014. – P. 740-755.",
}


# ------------------------------------------------------------------ сборка ----------------------

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def run(par, text, *, bold=False, italic=False, size=28, link=False):
    r = par.add_run(text)
    rpr = r._r.get_or_add_rPr()
    fonts = el("w:rFonts", w_ascii=FONT, w_hAnsi=FONT, w_cs=FONT, w_eastAsia=FONT)
    rpr.insert(0, fonts)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    r.font.size = Pt(size / 2)
    return r


def set_ppr(par, *, jc=None, first=None, after=0, line=240, keep_next=False):
    p = par._p.get_or_add_pPr()
    for child in list(p):
        p.remove(child)
    if keep_next:
        p.append(el("w:keepNext"))
    p.append(el("w:spacing", w_before=0, w_after=after, w_line=line, w_lineRule="auto"))
    if first is not None:
        p.append(el("w:ind", w_firstLine=first))
    if jc:
        p.append(el("w:jc", w_val=jc))


class Article:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        body = self.doc.element.body
        sect = body[-1]
        for child in list(body):
            if child is not sect:
                body.remove(child)      # sectPr остаётся: поля и размер страницы как в образце
        self.cites: dict[str, int] = {}

    # --- ссылки на источники
    def resolve(self, text: str) -> str:
        def sub(m):
            keys = [k.strip().lstrip("@") for k in m.group(1).split(",")]
            nums = []
            for k in keys:
                if k not in self.cites:
                    self.cites[k] = len(self.cites) + 1
                nums.append(self.cites[k])
            return "[" + ", ".join(str(n) for n in sorted(set(nums))) + "]"
        return re.sub(r"\[((?:@\w+(?:\s*,\s*)?)+)\]", sub, text)

    def blank(self, jc="both", first=INDENT, keep_next=False):
        set_ppr(self.doc.add_paragraph(), jc=jc, first=first, keep_next=keep_next)

    def line(self, text, *, jc, italic=False, bold=False, first=None, after=0, keep_next=False):
        par = self.doc.add_paragraph()
        set_ppr(par, jc=jc, first=first, after=after, keep_next=keep_next)
        run(par, text, bold=bold, italic=italic)
        return par

    def udk(self):
        par = self.doc.add_paragraph()
        p = par._p.get_or_add_pPr()
        run(par, UDK, bold=True)

    def email_line(self):
        par = self.doc.add_paragraph()
        set_ppr(par, jc="right")
        run(par, "e-mail:", italic=True)
        rid = self.doc.part.relate_to("mailto:" + EMAIL, RT.HYPERLINK, is_external=True)
        link = el("w:hyperlink", r_id=rid, w_history=1)
        r = el("w:r")
        rpr = el("w:rPr")
        rpr.append(el("w:rStyle", w_val="a5"))
        rpr.append(el("w:rFonts", w_ascii=FONT, w_hAnsi=FONT, w_cs=FONT))
        rpr.append(el("w:i"))
        rpr.append(el("w:iCs"))
        rpr.append(el("w:sz", w_val=28))
        rpr.append(el("w:szCs", w_val=28))
        r.append(rpr)
        t = el("w:t")
        t.text = EMAIL
        r.append(t)
        link.append(r)
        par._p.append(link)

    def header_block(self, authors, title, annot_label, annot, kw_label, kw, ru=True):
        for a in authors:
            if a == "EMAIL":
                self.email_line()
            else:
                self.line(a, jc="right", italic=True)
        self.line(title, jc="center", bold=True, italic=True)
        self.line(annot_label + " " + annot, jc="both", first=INDENT)
        self.line(kw_label + " " + kw, jc="both", first=INDENT)

    def heading(self, text):
        self.line(text, jc=None, bold=True, first=INDENT, keep_next=True)

    def paragraph(self, text):
        par = self.doc.add_paragraph()
        set_ppr(par, jc="both", first=INDENT)
        for i, chunk in enumerate(re.split(r"\^\{(.+?)\}", self.resolve(text))):
            r = run(par, chunk)
            if i % 2:                       # фрагмент в ^{...} набирается верхним индексом
                r.font.superscript = True

    def table(self, key):
        spec = TABLES[key]
        self.blank(keep_next=True)
        self.line(spec["caption"], jc="both", keep_next=True)
        self.blank(keep_next=True)
        rows, widths = spec["rows"], spec["widths"]
        t = self.doc.add_table(rows=len(rows), cols=len(widths))
        t.style = self.doc.styles["Table Grid"]
        tbl = t._tbl
        tblpr = tbl.tblPr
        for old in tblpr.findall(qn("w:tblW")):
            tblpr.remove(old)
        tblpr.append(el("w:tblW", w_w=sum(widths), w_type="dxa"))
        grid = tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate(rows):
            trpr = t.rows[ri]._tr.get_or_add_trPr()
            trpr.append(el("w:cantSplit"))
            for ci, text in enumerate(row):
                cell = t.cell(ri, ci)
                tcpr = cell._tc.get_or_add_tcPr()
                for old in tcpr.findall(qn("w:tcW")):
                    tcpr.remove(old)
                tcpr.insert(0, el("w:tcW", w_w=widths[ci], w_type="dxa"))
                par = cell.paragraphs[0]
                set_ppr(par, jc="center", keep_next=ri < len(rows) - 1)  # таблица не разрывается между страницами
                run(par, text, bold=(ri == 0), size=24)
        self.blank()

    def bibliography(self):
        self.blank(first=INDENT)
        self.line("Библиографический список", jc="both", bold=True, first=INDENT, after=0)
        for key, num in sorted(self.cites.items(), key=lambda kv: kv[1]):
            text = BIB.get(key) or SOURCES[key].replace("21.10.2026", ACCESS_DATE)
            text = re.sub(r" : ", ": ", text)
            # строки с длинными адресами при выравнивании по ширине растягиваются, поэтому для них выравнивание слева
            self.line(f"{num}. {text}", jc="left" if "URL:" in text else "both", first=INDENT)

    def build(self, out: Path):
        self.udk()
        self.header_block(AUTHOR_RU, TITLE_RU, "Аннотация.", ANNOT_RU, "Ключевые слова:", KEYWORDS_RU)
        self.header_block(AUTHOR_EN, TITLE_EN, "Annotation.", ANNOT_EN, "Keywords:", KEYWORDS_EN, ru=False)
        self.blank()
        for item in BODY:
            kind = item[0]
            if kind == "h":
                self.heading(item[1])
            elif kind == "p":
                self.paragraph(item[1])
            elif kind == "table":
                self.table(item[1])
        self.bibliography()
        cp = self.doc.core_properties          # в образце остались данные его автора
        cp.author = cp.last_modified_by = "А.Ю. Гридасов"
        cp.title = TITLE_RU
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    a = ap.parse_args()
    art = Article(a.template)
    art.build(a.out)
    print("источников:", len(art.cites), art.cites)
    print("готово:", a.out)

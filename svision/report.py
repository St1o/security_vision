"""Акт о нарушениях требований охраны труда (Word) за выбранную смену."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Callable, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .shifts import Shift
from .storage import ViolationRecord

FONT = "Times New Roman"


def _set_font(style_or_run, size: int | None = None, bold: bool | None = None) -> None:
    font = style_or_run.font
    font.name = FONT
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    element = getattr(style_or_run, "element", None)
    if element is None:
        element = style_or_run._element
    rpr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element.rPr
    rpr.rFonts.set(qn("w:eastAsia"), FONT)


def _paragraph(doc, text: str = "", bold: bool = False, align=None, size: int = 12):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    _set_font(run, size, bold)
    return p


def build_act(path: Path, records: Sequence[ViolationRecord], day: date, shift: Shift,
              camera: str, object_name: str, legal_basis: str,
              resolve_snapshot: Callable[[ViolationRecord], Path | None]) -> Path:
    """Формирует .docx и возвращает путь к нему."""
    doc = Document()
    _set_font(doc.styles["Normal"], 12)
    section = doc.sections[0]
    section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)

    number = f"{day:%Y%m%d}-{shift.number}"
    center = WD_ALIGN_PARAGRAPH.CENTER
    _paragraph(doc, f"АКТ № {number}", bold=True, align=center, size=14)
    _paragraph(doc, "о выявленных нарушениях требований охраны труда\nпри применении средств индивидуальной защиты",
               bold=True, align=center)

    _paragraph(doc, f"Объект: {object_name}")
    _paragraph(doc, f"Источник видеоконтроля: {camera}")
    _paragraph(doc, f"Дата: {day:%d.%m.%Y}")
    _paragraph(doc, f"Смена: {shift.name} ({shift.label})")
    _paragraph(doc, "Нарушения выявлены автоматизированной системой видеоконтроля «Security Vision». "
                    "Данные требуют проверки ответственным за охрану труда: система не заменяет контроль человека.")
    if legal_basis:
        _paragraph(doc, f"Нормативная база: {legal_basis}")

    if any(r.is_demo for r in records):
        _paragraph(doc, "ВНИМАНИЕ: часть записей помечена как ТЕСТОВЫЕ (демонстрационные данные) и "
                        "не является результатом реальных наблюдений.", bold=True)

    headers = ["№", "Время", "Зона", "Нарушение", "ID на видео", "Длит., с", "Примечание"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        _set_font(cell.paragraphs[0].add_run(text), 10, True)
    for i, rec in enumerate(records, 1):
        values = [
            str(i), rec.ts.strftime("%H:%M:%S"), rec.zone, rec.label,
            str(rec.track_id) if rec.track_id is not None else "—",
            f"{rec.duration_s:.0f}" if rec.duration_s is not None else "—",
            "тестовая запись" if rec.is_demo else "",
        ]
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = ""
            _set_font(cell.paragraphs[0].add_run(text), 10)

    widths = (Cm(1.0), Cm(2.0), Cm(2.2), Cm(4.6), Cm(2.0), Cm(1.8), Cm(3.0))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    photos = [(i, rec, resolve_snapshot(rec)) for i, rec in enumerate(records, 1)]
    photos = [(i, rec, p) for i, rec, p in photos if p is not None]
    if photos:
        _paragraph(doc, "")
        _paragraph(doc, "Фотофиксация", bold=True)
        for i, rec, photo in photos:
            doc.add_picture(str(photo), width=Cm(11))
            _paragraph(doc, f"Рис. {i}. {rec.label}, {rec.ts:%H:%M:%S}, {rec.zone}", size=10)

    _paragraph(doc, "")
    _paragraph(doc, "Ответственный за охрану труда: ______________ / ____________________ /")
    _paragraph(doc, "Нарушитель (заполняется после установления личности): ____________________________")
    _paragraph(doc, "С актом ознакомлен: ______________ / ____________________ /   «___» __________ 20___ г.")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
